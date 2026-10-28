"""Gera a planilha Excel com os lançamentos da pessoa."""

from io import BytesIO

import pandas as pd
from openpyxl.utils import get_column_letter

KIND_LABELS = {"expense": "Gasto", "income": "Ganho"}


def transactions_to_excel(transactions: list[dict]) -> bytes:
    """Uma linha por lançamento, com data, tipo, categoria, descrição e valor em reais."""
    df = pd.DataFrame(
        {
            "Data": [t["occurred_on"] for t in transactions],
            "Tipo": [KIND_LABELS[t["kind"]] for t in transactions],
            "Categoria": [t["category"] for t in transactions],
            "Descrição": [t["description"] or "" for t in transactions],
            # Gastos negativos e ganhos positivos: assim um SOMA() no Excel dá o saldo
            "Valor (R$)": [
                (-1 if t["kind"] == "expense" else 1) * t["amount_cents"] / 100 for t in transactions
            ],
        }
    )

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Lançamentos", index=False)
        sheet = writer.sheets["Lançamentos"]
        sheet.freeze_panes = "A2"  # o cabeçalho fica parado ao rolar
        for row in sheet.iter_rows(min_row=2):
            row[0].number_format = "DD/MM/YYYY"
            row[4].number_format = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
        for i, width in enumerate([12, 8, 18, 32, 14], start=1):
            sheet.column_dimensions[get_column_letter(i)].width = width
    return buffer.getvalue()
